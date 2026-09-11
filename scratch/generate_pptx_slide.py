import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_feasibility_slide():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # 1. Background Fill (Dark Teal Orbital Theme)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(6, 31, 35) # #061F23
    bg.line.fill.background() # No border
    
    # Palette definition
    c_white = RGBColor(255, 255, 255)
    c_mint = RGBColor(133, 239, 208)   # #85EFD0
    c_emerald = RGBColor(16, 185, 129) # #10B981
    c_card_bg = RGBColor(11, 46, 51)   # #0B2E33
    c_card_border = RGBColor(26, 77, 84) # #1A4D54
    c_highlight_bg = RGBColor(13, 64, 58) # #0D403A
    c_gray_text = RGBColor(160, 185, 182) # #A0B9B6
    c_neutral_bar = RGBColor(90, 115, 118) # #5A7376
    
    # -------------------------------------------------------------
    # HEADER ZONE
    # -------------------------------------------------------------
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.8))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = "SatQuery AI — Feasibility at a glance"
    p0.font.name = "Arial"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = c_white
    
    p1 = tf.add_paragraph()
    p1.text = "Pre-Implementation Architecture Targets & Comparative Feasibility Projection"
    p1.font.name = "Arial"
    p1.font.size = Pt(11)
    p1.font.color.rgb = c_mint
    
    # -------------------------------------------------------------
    # TOP ROW: 6 SCORE CARDS
    # -------------------------------------------------------------
    cards_data = [
        ("Technical Feasibility", "9.5", False),
        ("Compute & Hardware", "10.0", False),
        ("RS & Data", "8.5", False),
        ("Economic Viability", "9.5", False),
        ("Market Fit", "9.0", False),
        ("Overall Target", "9.3", True) # Highlighted
    ]
    
    card_width = Inches(1.82)
    card_height = Inches(0.88)
    card_gap = Inches(0.16)
    start_x = Inches(0.8)
    cards_y = Inches(1.22)
    
    for idx, (label, score, is_highlight) in enumerate(cards_data):
        cx = start_x + idx * (card_width + card_gap)
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cards_y, card_width, card_height)
        c_shape.fill.solid()
        if is_highlight:
            c_shape.fill.fore_color.rgb = c_highlight_bg
            c_shape.line.color.rgb = c_emerald
            c_shape.line.width = Pt(2.0)
        else:
            c_shape.fill.fore_color.rgb = c_card_bg
            c_shape.line.color.rgb = c_card_border
            c_shape.line.width = Pt(1.0)
            
        ctf = c_shape.text_frame
        ctf.word_wrap = True
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ctf.margin_left = ctf.margin_right = Inches(0.08)
        ctf.margin_top = ctf.margin_bottom = Inches(0.04)
        
        lp = ctf.paragraphs[0]
        lp.text = label.upper()
        lp.alignment = PP_ALIGN.CENTER
        lp.font.name = "Arial"
        lp.font.size = Pt(7.5)
        lp.font.bold = True
        lp.font.color.rgb = c_emerald if is_highlight else c_gray_text
        
        sp = ctf.add_paragraph()
        sp.text = f"{score} / 10"
        sp.alignment = PP_ALIGN.CENTER
        sp.font.name = "Arial"
        sp.font.size = Pt(14)
        sp.font.bold = True
        sp.font.color.rgb = c_mint if is_highlight else c_white
        
        tp = ctf.add_paragraph()
        tp.text = "TARGET PROJECTION" if not is_highlight else "SYSTEM GOAL"
        tp.alignment = PP_ALIGN.CENTER
        tp.font.name = "Arial"
        tp.font.size = Pt(6)
        tp.font.color.rgb = c_emerald if is_highlight else RGBColor(100, 135, 138)

    # -------------------------------------------------------------
    # MIDDLE ZONE: LEFT (HORIZONTAL STACKED BAR) & RIGHT (VERTICAL BARS)
    # -------------------------------------------------------------
    mid_y = Inches(2.28)
    mid_w = Inches(5.72)
    mid_h = Inches(2.35)
    
    # --- MIDDLE-LEFT CONTAINER ---
    ml_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), mid_y, mid_w, mid_h)
    ml_box.fill.solid()
    ml_box.fill.fore_color.rgb = c_card_bg
    ml_box.line.color.rgb = c_card_border
    ml_box.line.width = Pt(1)
    
    ml_tf = ml_box.text_frame
    ml_tf.margin_top = Inches(0.12)
    ml_tf.margin_left = Inches(0.18)
    mlp0 = ml_tf.paragraphs[0]
    mlp0.text = "Projected VRAM budget (6GB target)"
    mlp0.font.name = "Arial"
    mlp0.font.size = Pt(11)
    mlp0.font.bold = True
    mlp0.font.color.rgb = c_white
    
    # Subtitle with cap
    mlp1 = ml_tf.add_paragraph()
    mlp1.text = "Total Memory Axis Capped at 5.64 GB (Usable RTX 4050 Ceiling)"
    mlp1.font.name = "Arial"
    mlp1.font.size = Pt(8)
    mlp1.font.color.rgb = c_mint
    
    # Stacked bar segments (Total 5.64GB)
    # Width of full bar = 5.3 inches
    bar_x = Inches(1.0)
    bar_y = Inches(2.95)
    full_bar_w = 5.32
    max_vram = 5.64
    
    segments = [
        ("Interpreter", 1.45, RGBColor(24, 110, 100), "1.45GB"),
        ("Vision model", 2.42, RGBColor(16, 185, 129), "2.42GB"),
        ("Workspace", 1.20, RGBColor(40, 145, 150), "1.20GB"),
        ("Headroom", 0.57, RGBColor(80, 180, 160), "0.57GB")
    ]
    
    cur_x = bar_x
    for name, gb, col, label in segments:
        seg_w = Inches(full_bar_w * (gb / max_vram))
        seg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cur_x, bar_y, seg_w, Inches(0.55))
        seg.fill.solid()
        seg.fill.fore_color.rgb = col
        seg.line.color.rgb = c_card_bg
        seg.line.width = Pt(1)
        
        stf = seg.text_frame
        stf.vertical_anchor = MSO_ANCHOR.MIDDLE
        stf.margin_top = stf.margin_bottom = stf.margin_left = stf.margin_right = 0
        stp = stf.paragraphs[0]
        stp.text = f"{label}"
        stp.alignment = PP_ALIGN.CENTER
        stp.font.name = "Arial"
        stp.font.size = Pt(8.5)
        stp.font.bold = True
        stp.font.color.rgb = c_white
        
        cur_x += seg_w
        
    # Legend below stacked bar
    leg_x = Inches(1.0)
    leg_y = Inches(3.6)
    leg_box = slide.shapes.add_textbox(leg_x, leg_y, Inches(5.32), Inches(0.45))
    ltf = leg_box.text_frame
    ltf.margin_top = ltf.margin_left = ltf.margin_right = ltf.margin_bottom = 0
    lp = ltf.paragraphs[0]
    lp.text = "■ Interpreter (1.45GB)    ■ Vision (2.42GB)    ■ Workspace (1.20GB)    ■ Safety Headroom (0.57GB)"
    lp.font.name = "Arial"
    lp.font.size = Pt(7.5)
    lp.font.color.rgb = c_gray_text
    
    # Caption underneath
    lp_cap = ltf.add_paragraph()
    lp_cap.text = "Caption: Projected from published model sizes — not yet a measured run."
    lp_cap.font.name = "Arial"
    lp_cap.font.size = Pt(8)
    lp_cap.font.italic = True
    lp_cap.font.color.rgb = c_mint
    
    # --- MIDDLE-RIGHT CONTAINER ---
    mr_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), mid_y, mid_w, mid_h)
    mr_box.fill.solid()
    mr_box.fill.fore_color.rgb = c_card_bg
    mr_box.line.color.rgb = c_card_border
    mr_box.line.width = Pt(1)
    
    mr_tf = mr_box.text_frame
    mr_tf.margin_top = Inches(0.12)
    mr_tf.margin_left = Inches(0.18)
    mrp0 = mr_tf.paragraphs[0]
    mrp0.text = "Projected monthly compute cost"
    mrp0.font.name = "Arial"
    mrp0.font.size = Pt(11)
    mrp0.font.bold = True
    mrp0.font.color.rgb = c_white
    
    mrp1 = mr_tf.add_paragraph()
    mrp1.text = "Continuous 24/7 Hosting Comparison (USD/month target projection)"
    mrp1.font.name = "Arial"
    mrp1.font.size = Pt(8)
    mrp1.font.color.rgb = c_mint
    
    # 4 Vertical Bars: A100 $2,642, Multi-API $1,350, Cloud micro-node $201, Local edge (proposed) $0
    vbars_data = [
        ("A100 server", 2642, "$2,642", c_neutral_bar, False),
        ("Multi-API", 1350, "$1,350", c_neutral_bar, False),
        ("Cloud micro", 201, "$201", c_neutral_bar, False),
        ("Local edge\n(proposed)", 0, "$0 Target", c_emerald, True)
    ]
    
    chart_base_y = Inches(4.18)
    max_h_val = 2800.0
    bar_area_h = 1.10 # inches
    vbar_w = Inches(0.92)
    vbar_start_x = Inches(7.1)
    vbar_gap = Inches(0.42)
    
    for v_idx, (v_name, cost, cost_lbl, v_color, is_accent) in enumerate(vbars_data):
        vx = vbar_start_x + v_idx * (vbar_w + vbar_gap)
        
        # Calculate bar height
        bar_h_in = Inches(max(0.04, (cost / max_h_val) * bar_area_h)) if cost > 0 else Inches(0.06)
        vy = chart_base_y - bar_h_in
        
        # Draw vertical bar
        vbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, vx, vy, vbar_w, bar_h_in)
        vbar.fill.solid()
        vbar.fill.fore_color.rgb = v_color
        vbar.line.fill.background()
        
        # Cost label above bar
        c_lbl_box = slide.shapes.add_textbox(vx - Inches(0.1), vy - Inches(0.28), vbar_w + Inches(0.2), Inches(0.26))
        ctf = c_lbl_box.text_frame
        ctf.margin_top = ctf.margin_bottom = ctf.margin_left = ctf.margin_right = 0
        cp = ctf.paragraphs[0]
        cp.text = cost_lbl
        cp.alignment = PP_ALIGN.CENTER
        cp.font.name = "Arial"
        cp.font.size = Pt(8.5)
        cp.font.bold = True
        cp.font.color.rgb = c_emerald if is_accent else c_white
        
        # Category label below bar
        n_lbl_box = slide.shapes.add_textbox(vx - Inches(0.1), chart_base_y + Inches(0.05), vbar_w + Inches(0.2), Inches(0.35))
        ntf = n_lbl_box.text_frame
        ntf.margin_top = ntf.margin_bottom = ntf.margin_left = ntf.margin_right = 0
        np_ = ntf.paragraphs[0]
        np_.text = v_name
        np_.alignment = PP_ALIGN.CENTER
        np_.font.name = "Arial"
        np_.font.size = Pt(7.5)
        np_.font.color.rgb = c_emerald if is_accent else c_gray_text

    # -------------------------------------------------------------
    # BOTTOM ZONE: 5-COLUMN COMPARISON TABLE
    # -------------------------------------------------------------
    table_x = Inches(0.8)
    table_y = Inches(4.78)
    table_w = Inches(11.733)
    table_h = Inches(2.28)
    
    rows_count = 6 # 1 header + 5 feature rows
    cols_count = 5
    
    table_shape = slide.shapes.add_table(rows_count, cols_count, table_x, table_y, table_w, table_h)
    tbl = table_shape.table
    
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(2.2) # SatQuery AI (proposed)
    tbl.columns[2].width = Inches(2.1) # GPT-4o / Gemini 1.5
    tbl.columns[3].width = Inches(2.1) # Prithvi-100M
    tbl.columns[4].width = Inches(2.133) # EarthGPT
    
    table_content = [
        ["Capability / Dimension", "SatQuery AI (proposed)", "GPT-4o / Gemini 1.5", "Prithvi-100M (NASA)", "EarthGPT"],
        ["Runs on 6GB laptop GPU", "Target: yes", "No (Cloud API)", "Backbone only", "No (≥24GB)"],
        ["Raw SAR ingestion", "Target: yes", "No", "Limited", "No"],
        ["Bitemporal change VQA", "Target: yes", "Unreliable", "Detection only", "Uncalibrated"],
        ["Explainable evidence", "Target: yes", "No", "No", "Limited"],
        ["Offline / air-gapped", "Target: yes", "No", "Yes", "No"]
    ]
    
    for r_idx, row in enumerate(table_content):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            
            # Styling cell fill
            cell.fill.solid()
            if r_idx == 0:
                if c_idx == 1:
                    cell.fill.fore_color.rgb = RGBColor(16, 80, 70) # Accent header
                else:
                    cell.fill.fore_color.rgb = RGBColor(9, 38, 43)
            else:
                if c_idx == 1:
                    cell.fill.fore_color.rgb = RGBColor(12, 50, 48) # Proposed column highlighted
                else:
                    cell.fill.fore_color.rgb = RGBColor(8, 32, 36) if r_idx % 2 == 0 else RGBColor(6, 26, 29)
                    
            # Text formatting
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
            p.font.name = "Arial"
            
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(8.5)
                p.font.color.rgb = c_emerald if c_idx == 1 else c_white
            else:
                p.font.size = Pt(8.0)
                if c_idx == 1:
                    p.font.bold = True
                    p.font.color.rgb = c_mint # "Target: yes" highlighted in mint
                else:
                    p.font.color.rgb = c_gray_text
                    
    # Save Presentation
    output_pptx = "docs/SatQuery_Feasibility_Slide.pptx"
    prs.save(output_pptx)
    print(f"Successfully generated PowerPoint presentation: {output_pptx}")

if __name__ == "__main__":
    create_feasibility_slide()
